"""
Ayudas para armar los documentos de Word (python-docx)
======================================================
Títulos numerados, ecuaciones editables de Word (LaTeX → MathML → OMML con el
conversor de Office), figuras y tablas con leyenda, índice y número de página.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from latex2mathml.converter import convert
from lxml import etree

XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
_XSLT = etree.XSLT(etree.parse(str(XSL)))
AZUL = RGBColor(0x1F, 0x4E, 0x8C)
GRIS = RGBColor(0x55, 0x60, 0x6B)


def _sombra(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    tc.append(s)


def _campo(parrafo, instruccion, texto="…"):
    r = parrafo.add_run()
    for tipo, extra in (("begin", None), ("instr", instruccion), ("separate", None), ("texto", texto), ("end", None)):
        if tipo == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = extra
            r._r.append(e)
        elif tipo == "texto":
            t = OxmlElement("w:t")
            t.text = extra
            r._r.append(t)
        else:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), tipo)
            r._r.append(f)


class Documento:
    def __init__(self, titulo, subtitulo, codigo, version="1.0", fecha="28 de octubre de 2026",
                 autor="Roldán Rodríguez — Ingeniería en Ciencias de la Computación e Inteligencia Artificial"):
        self.d = Document()
        self.n = {"h1": 0, "h2": 0, "h3": 0, "fig": 0, "tab": 0, "ec": 0}
        self.titulo = titulo
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)          # carta
        for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, lado, Cm(2.5))
        self.ancho = s.page_width - s.left_margin - s.right_margin
        estilos = self.d.styles
        n = estilos["Normal"]
        n.font.name, n.font.size = "Calibri", Pt(11)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.12
        for nivel, tam in ((1, 16), (2, 13), (3, 11.5)):
            h = estilos[f"Heading {nivel}"]
            h.font.name, h.font.size, h.font.color.rgb, h.font.bold = "Calibri", Pt(tam), AZUL, True
            h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")
            h.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        if "Leyenda" not in [x.name for x in estilos]:
            ley = estilos.add_style("Leyenda", 1)
            ley.base_style = n
            ley.font.size, ley.font.italic, ley.font.color.rgb = Pt(9.5), True, GRIS
            ley.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ley.paragraph_format.space_after = Pt(10)
        # encabezado y pie
        enc = s.header.paragraphs[0]
        enc.text = f"{titulo} · {codigo} · v{version}"
        enc.runs[0].font.size, enc.runs[0].font.color.rgb = Pt(8.5), GRIS
        pie = s.footer.paragraphs[0]
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = pie.add_run("Página ")
        r.font.size = Pt(9)
        _campo(pie, "PAGE", "1")
        self._portada(titulo, subtitulo, codigo, version, fecha, autor)

    # ---------- estructura ----------
    def _portada(self, titulo, subtitulo, codigo, version, fecha, autor):
        for _ in range(5):
            self.d.add_paragraph()
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("PRÓTESIS MIOELÉCTRICA CON RECONOCIMIENTO DE GESTOS POR IA")
        r.font.size, r.font.bold, r.font.color.rgb = Pt(12), True, GRIS
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(titulo)
        r.font.size, r.font.bold, r.font.color.rgb = Pt(26), True, AZUL
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitulo)
        r.font.size, r.font.color.rgb = Pt(13), GRIS
        for _ in range(8):
            self.d.add_paragraph()
        t = self.d.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in (("Autor", autor), ("Código del documento", codigo), ("Versión", version), ("Fecha", fecha),
                     ("Repositorio", "github.com/roldan0411/mano")):
            c = t.add_row().cells
            c[0].text, c[1].text = k, v
            c[0].paragraphs[0].runs[0].font.bold = True
            c[0].width, c[1].width = Cm(4.5), Cm(11)
        self.salto()

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def indice(self):
        t = self.d.add_paragraph()
        r = t.add_run("Contenido")
        r.font.size, r.font.bold, r.font.color.rgb = Pt(16), True, AZUL
        _campo(self.d.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Actualizar el índice: clic derecho → Actualizar campo.")
        self.salto()

    def h1(self, texto, numerar=True):
        if numerar:
            self.n["h1"] += 1
            self.n["h2"] = self.n["h3"] = 0
            texto = f"{self.n['h1']}. {texto}"
        return self.d.add_heading(texto, 1)

    def h2(self, texto):
        self.n["h2"] += 1
        self.n["h3"] = 0
        return self.d.add_heading(f"{self.n['h1']}.{self.n['h2']} {texto}", 2)

    def h3(self, texto):
        self.n["h3"] += 1
        return self.d.add_heading(f"{self.n['h1']}.{self.n['h2']}.{self.n['h3']} {texto}", 3)

    # ---------- texto ----------
    def _runs(self, p, texto):
        """**negrita** y *cursiva* simples."""
        for parte in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", texto):
            if not parte:
                continue
            if parte.startswith("**"):
                p.add_run(parte[2:-2]).bold = True
            elif parte.startswith("*"):
                p.add_run(parte[1:-1]).italic = True
            else:
                p.add_run(parte)
        return p

    def p(self, texto, estilo=None):
        par = self.d.add_paragraph(style=estilo)
        par.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return self._runs(par, texto)

    def vinetas(self, items, numeradas=False):
        for it in items:
            self._runs(self.d.add_paragraph(style="List Number" if numeradas else "List Bullet"), it)

    def nota(self, texto):
        par = self.p(texto)
        for r in par.runs:
            r.font.size, r.font.color.rgb = Pt(9.5), GRIS
        return par

    # ---------- ecuaciones ----------
    def ec(self, latex, donde=None):
        """Ecuación editable de Word, centrada y numerada. donde: lista de (símbolo, significado)."""
        self.n["ec"] += 1
        par = self.d.add_paragraph()
        pf = par.paragraph_format
        pf.tab_stops.add_tab_stop(int(self.ancho / 2), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(self.ancho, WD_TAB_ALIGNMENT.RIGHT)
        par.add_run("\t")
        omml = _XSLT(etree.fromstring(convert(latex).encode()))
        par._p.append(omml.getroot())
        par.add_run(f"\t({self.n['ec']})")
        if donde:
            q = self.d.add_paragraph()
            q.paragraph_format.left_indent = Cm(1)
            q.add_run("donde ").italic = True
            for k, (sim, sig) in enumerate(donde):
                om = _XSLT(etree.fromstring(convert(sim).encode()))
                q._p.append(om.getroot())
                q.add_run(f" {sig}" + ("; " if k < len(donde) - 1 else "."))
            for r in q.runs:
                r.font.size = Pt(10)
        return self.n["ec"]

    # ---------- figuras y tablas ----------
    def fig(self, ruta, leyenda, ancho_cm=16):
        self.n["fig"] += 1
        par = self.d.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(str(ruta), width=Cm(min(ancho_cm, 16.5)))
        self._runs(self.d.add_paragraph(style="Leyenda"), f"Figura {self.n['fig']}. {leyenda}")
        return self.n["fig"]

    def tabla(self, encabezados, filas, leyenda=None, anchos=None, tam=9.5, centrar=None):
        if leyenda:
            self.n["tab"] += 1
            par = self._runs(self.d.add_paragraph(style="Leyenda"), f"Tabla {self.n['tab']}. {leyenda}")
            par.paragraph_format.keep_with_next = True
            par.paragraph_format.space_after = Pt(3)
        t = self.d.add_table(rows=1, cols=len(encabezados))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, e in enumerate(encabezados):
            c = t.rows[0].cells[k]
            c.text = ""
            r = c.paragraphs[0].add_run(str(e))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(tam), RGBColor(255, 255, 255)
            _sombra(c, "1F4E8C")
        for i, fila in enumerate(filas):
            celdas = t.add_row().cells
            for k, v in enumerate(fila):
                celdas[k].text = ""
                self._runs(celdas[k].paragraphs[0], str(v))
                for r in celdas[k].paragraphs[0].runs:
                    r.font.size = Pt(tam)
                if centrar and k in centrar:
                    celdas[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if i % 2:
                    _sombra(celdas[k], "EEF3F9")
        if anchos:
            for fila in t.rows:
                for k, w in enumerate(anchos):
                    fila.cells[k].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def ficha(self, pares, anchos=(4.2, 12.3)):
        t = self.d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in pares:
            c = t.add_row().cells
            c[0].text = ""
            r = c[0].paragraphs[0].add_run(k)
            r.bold, r.font.size = True, Pt(9.5)
            _sombra(c[0], "E3EBF5")
            c[1].text = ""
            self._runs(c[1].paragraphs[0], str(v))
            for r in c[1].paragraphs[0].runs:
                r.font.size = Pt(9.5)
            c[0].width, c[1].width = Cm(anchos[0]), Cm(anchos[1])
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def guardar(self, ruta):
        Path(ruta).parent.mkdir(parents=True, exist_ok=True)
        self.d.save(ruta)
        return ruta
